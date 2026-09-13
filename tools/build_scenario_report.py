#!/usr/bin/env python3
"""Build the NAVIS scenario report from the repository Markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "16324F"
BLUE = "24577A"
TEAL = "2E7D7B"
LIGHT_BLUE = "EAF1F6"
LIGHT_GREY = "F5F7F9"
MID_GREY = "D9D9D9"
DARK_GREY = RGBColor(55, 65, 75)
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GREY, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_heading(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Pagina ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = DARK_GREY
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_inline_runs(paragraph, text: str, *, color=BLACK, size=None) -> None:
    """Render the small subset of Markdown used by the source document."""
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            run.font.color.rgb = color
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt((size or 10.2) - 0.4)
            run.font.color.rgb = RGBColor(30, 78, 99)
        if size and token.startswith("**"):
            run.font.size = Pt(size)
        if token.startswith("**"):
            run.font.color.rgb = color
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)


def clean_markdown(text: str) -> str:
    return re.sub(r"\*\*|`", "", text).replace("  ", " ").strip()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.77)
    section.right_margin = Inches(0.77)
    section.header_distance = Inches(0.33)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.1)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(8)

    for style_name, size, space_before, color in (
        ("Heading 1", 18, 14, BLACK),
        ("Heading 2", 14, 12, BLACK),
        ("Heading 3", 11.5, 9, BLACK),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(10)
    styles["List Bullet"].paragraph_format.space_after = Pt(2.5)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(10)
    styles["List Number"].paragraph_format.space_after = Pt(2.5)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "NAVIS CYBER READINESS   ·   DOCUMENTAZIONE SCENARI"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(36, 87, 122)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc: Document, lines: list[str]) -> int:
    title_text = lines[0][2:].strip()
    subtitle_text = lines[2][3:].strip()

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag_run = tag.add_run("MARITIME CYBER AWARENESS · SISTEMA LOCALE")
    tag_run.bold = True
    tag_run.font.name = "Aptos"
    tag_run.font.size = Pt(9)
    tag_run.font.color.rgb = RGBColor(36, 87, 122)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(title_text)
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    sr = subtitle.add_run(subtitle_text)
    sr.font.name = "Aptos Display"
    sr.font.size = Pt(17)
    sr.font.color.rgb = DARK_GREY

    meta = []
    index = 4
    while index < len(lines) and lines[index].strip():
        meta.append(clean_markdown(lines[index]))
        index += 1

    meta_table = doc.add_table(rows=len(meta), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    labels = []
    values = []
    for item in meta:
        if ":" in item:
            label, value = item.split(":", 1)
        else:
            label, value = "", item
        labels.append(label.strip())
        values.append(value.strip())
    for row_index, row in enumerate(meta_table.rows):
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(4.2)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, 100, 130, 100, 130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(labels[row_index].upper())
        r0.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = RGBColor(36, 87, 122)
        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(values[row_index])
        r1.font.size = Pt(9.5)
    set_table_borders(meta_table)

    while index < len(lines) and not lines[index].strip():
        index += 1
    if index < len(lines) and lines[index].startswith(">"):
        disclaimer = lines[index][1:].strip()
        callout = doc.add_table(rows=1, cols=1)
        callout.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(callout, NAVY, "4")
        cell = callout.cell(0, 0)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 150, 175, 150, 175)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline_runs(p, disclaimer, color=RGBColor(255, 255, 255), size=9)
        index += 1

    doc.add_page_break()
    return index


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    font_size = 7.4 if column_count >= 6 else 8.2 if column_count >= 4 else 9.0
    for row_index, source_row in enumerate(rows):
        row = table.rows[row_index]
        for col_index in range(column_count):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 75, 80, 75, 80)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                color = RGBColor(255, 255, 255)
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, LIGHT_GREY)
                color = BLACK
            value = source_row[col_index] if col_index < len(source_row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, value, color=color, size=font_size)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "B7CEDD", "5")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 120, 150, 120, 150)
    p = cell.paragraphs[0]
    add_inline_runs(p, text, color=BLACK, size=9.3)


def should_break_before_heading(text: str, level: int) -> bool:
    if level == 2 and text.startswith("Scenario "):
        return not text.startswith("Scenario 1 ")
    if level == 1 and re.match(r"(?:4|5|6)\.", text):
        return True
    return False


def add_body(doc: Document, lines: list[str], start: int) -> None:
    index = start
    paragraph_buffer: list[tuple[str, bool]] = []
    in_code = False
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        buffered_lines = list(paragraph_buffer)
        paragraph_buffer.clear()
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        for line_index, (text_value, hard_break) in enumerate(buffered_lines):
            add_inline_runs(p, text_value)
            if hard_break:
                p.add_run().add_break()
            elif line_index < len(buffered_lines) - 1:
                p.add_run(" ")

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                code_table = doc.add_table(rows=1, cols=1)
                set_table_borders(code_table, "C7D2DB", "5")
                cell = code_table.cell(0, 0)
                set_cell_shading(cell, LIGHT_GREY)
                set_cell_margins(cell, 115, 145, 115, 145)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Aptos Mono"
                run.font.size = Pt(8.8)
                in_code = False
                code_lines = []
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            hashes, heading_text = heading_match.groups()
            level = len(hashes) - 1
            if should_break_before_heading(heading_text, level):
                doc.add_page_break()
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(p, heading_text)
            set_repeat_heading(p)
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            add_callout(doc, stripped[1:].strip())
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet_match.group(1))
            index += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if number_match:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.27)
            p.paragraph_format.first_line_indent = Inches(-0.27)
            p.paragraph_format.space_after = Pt(2.5)
            number_run = p.add_run(f"{number_match.group(1)}.\t")
            number_run.font.name = "Aptos"
            number_run.font.size = Pt(10)
            add_inline_runs(p, number_match.group(2))
            index += 1
            continue

        paragraph_buffer.append((stripped, raw.endswith("  ")))
        index += 1

    flush_paragraph()


def set_core_properties(doc: Document) -> None:
    properties = doc.core_properties
    properties.title = "NAVIS Cyber Readiness — Report dettagliato degli scenari"
    properties.subject = "Documentazione formativa degli scenari di social engineering"
    properties.author = "NAVIS Cyber Readiness"
    properties.keywords = "cyber awareness, marina, social engineering, formazione, scenari"
    properties.comments = "Documento didattico; scenari interamente simulati."


def build(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise ValueError("Il file Markdown deve iniziare con un titolo di primo livello")

    doc = Document()
    configure_document(doc)
    set_core_properties(doc)
    body_start = add_cover(doc, lines)
    add_body(doc, lines, body_start)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    if len(sys.argv) != 3:
        print("Uso: build_scenario_report.py SOURCE.md OUTPUT.docx", file=sys.stderr)
        return 2
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    build(source, output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
